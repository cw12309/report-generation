"""Report 工作台 API 路由"""

import asyncio
import json
import logging
from pathlib import Path
from typing import Optional

from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse

from .schemas import (
    GenerateRequest,
    IterateRequest,
    SaveRequest,
    ReportResponse,
    ReportSession,
)
from .service import ReportService, CHAT_STANDARD_REPLY

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/api/report-workbench", tags=["report-workbench"])

# 全局服务实例
_service = ReportService()

# 简单的内存存储（后续可改为数据库）
_sessions: dict[str, ReportSession] = {}

# 历史报告持久化（JSON 文件）
_HISTORY_FILE = Path(__file__).resolve().parents[2] / "data" / "report_history.json"
_HISTORY_LIMIT = 100


def _load_history() -> list:
    if not _HISTORY_FILE.exists():
        return []
    try:
        with open(_HISTORY_FILE, "r", encoding="utf-8") as f:
            data = json.load(f)
        return data if isinstance(data, list) else []
    except Exception as e:
        logger.warning(f"读取报告历史失败: {e}")
        return []


def _write_history(history: list):
    try:
        _HISTORY_FILE.parent.mkdir(parents=True, exist_ok=True)
        with open(_HISTORY_FILE, "w", encoding="utf-8") as f:
            json.dump(history, f, ensure_ascii=False, indent=2)
    except Exception as e:
        logger.warning(f"写入报告历史失败: {e}")


def _persist_history(session: ReportSession):
    """将会话写入历史（按 conversation_id upsert，保留最近 N 条）"""
    history = _load_history()
    history = [h for h in history if h.get("conversation_id") != session.conversation_id]
    history.insert(0, session.dict())
    _write_history(history[:_HISTORY_LIMIT])


def _get_or_restore_session(conversation_id: str) -> Optional[ReportSession]:
    """获取会话；内存中不存在时从历史文件恢复（服务重启后历史报告仍可对话修改）"""
    session = _sessions.get(conversation_id)
    if session:
        return session
    for h in _load_history():
        if h.get("conversation_id") == conversation_id:
            try:
                restored = ReportSession(**h)
                _sessions[conversation_id] = restored
                logger.info(f"从历史恢复会话: {conversation_id}")
                return restored
            except Exception as e:
                logger.warning(f"恢复历史会话失败 {conversation_id}: {e}")
                return None
    return None


def _sse(data: dict) -> str:
    """序列化为 SSE data 帧"""
    return f"data: {json.dumps(data, ensure_ascii=False)}\n\n"


# 流式队列中允许下发的事件类型（token=报告正文增量，reasoning=思考过程增量，notice=系统提示）
_STREAM_EVENT_TYPES = ("token", "reasoning", "notice")


def _make_stream_event(kind: str, payload: str) -> str:
    """把队列中的 (kind, payload) 转成 SSE 帧；未知类型返回空串（调用方跳过）"""
    if kind in _STREAM_EVENT_TYPES and payload:
        return _sse({"type": kind, "content": payload})
    return ""


async def _emit_sse_event(event_type: str, data: dict):
    """SSE事件发射器（占位）"""
    logger.debug(f"SSE事件: {event_type} = {data}")


@router.get("/history")
async def list_history():
    """历史报告列表（摘要，按更新时间倒序）"""
    history = _load_history()
    items = []
    for h in history:
        report = h.get("report") or {}
        findings = report.get("imaging_findings") or ""
        items.append({
            "conversation_id": h.get("conversation_id"),
            "report_id": h.get("report_id"),
            "route_id": h.get("route_id"),
            "route_name": h.get("route_name"),
            "question_round": h.get("question_round", 0),
            "current_version": h.get("current_version", 0),
            "created_at": h.get("created_at"),
            "updated_at": h.get("updated_at"),
            "preview": findings[:60].replace("\n", " "),
        })
    return {"history": items}


@router.get("/history/{conversation_id}")
async def get_history_item(conversation_id: str):
    """获取历史报告完整内容"""
    for h in _load_history():
        if h.get("conversation_id") == conversation_id:
            return {"session": h}
    raise HTTPException(status_code=404, detail="历史报告不存在")


@router.delete("/history/{conversation_id}")
async def delete_history_item(conversation_id: str):
    """删除历史报告"""
    history = _load_history()
    new_history = [h for h in history if h.get("conversation_id") != conversation_id]
    if len(new_history) == len(history):
        raise HTTPException(status_code=404, detail="历史报告不存在")
    _write_history(new_history)
    _sessions.pop(conversation_id, None)
    return {"success": True}


@router.get("/models")
async def list_models():
    """列出可选模型（config.yml 中有效的 LLM 条目）"""
    from app.config import get_llm_entries, get_llm_model, get_llm_base_url
    entries = get_llm_entries()
    active_model = get_llm_model()
    active_url = get_llm_base_url()
    active_name = ""
    for e in entries:
        if e["model"] == active_model and e["base_url"] == active_url:
            active_name = e["name"]
            break
    return {"models": entries, "active": active_name or (entries[0]["name"] if entries else "")}


@router.get("/models/check")
async def check_model(name: str = ""):
    """探活指定模型是否可响应（选择模型前预校验）"""
    from app.chat.llm_client import check_model_available
    ok, message = await check_model_available(name)
    return {"ok": ok, "message": message}


@router.get("/models/thinking-test")
async def test_model_thinking(name: str = ""):
    """实测指定模型的思考开关效果（开/关各发一次请求比较输出）"""
    from app.chat.llm_client import test_thinking_switch
    return await test_thinking_switch(name)


@router.get("/sessions/{conversation_id}/generate")
async def generate_report(
    conversation_id: str,
    input_text: str = "",
    language: str = "zh-CN",
    modality: str = "",
    body_part: str = "",
    route_id: str = "",
    llm_name: str = "",
    enable_thinking: bool = False
):
    """生成报告（SSE流式，GET方法用于EventSource）"""
    
    # 构建请求对象
    request = GenerateRequest(
        input_text=input_text,
        language=language,
        modality=modality or None,
        body_part=body_part or None,
        route_id=route_id or None,
        llm_name=llm_name or "",
        enable_thinking=bool(enable_thinking)
    )
    
    async def event_stream():
        """SSE事件流（真流式：token 实时推送 + 完整耗时记录）"""
        # 会话ID
        yield _sse({'type': 'conversation', 'conversation_id': conversation_id})

        # 闲聊分流：非报告类输入返回标准回复，不调用LLM
        if not _service.is_report_intent(input_text):
            logger.info(f"非报告类输入，走闲聊标准回复: {input_text[:30]}")
            yield _sse({'type': 'chat', 'content': CHAT_STANDARD_REPLY})
            yield _sse({'type': 'done', 'status': 'chat', 'elapsed_seconds': 0.0})
            return

        queue: asyncio.Queue = asyncio.Queue()
        loop = asyncio.get_event_loop()

        # 提前解析路由并立即下发，减少用户等待空白期
        early_route = _service.route_resolver.resolve(
            input_text=input_text,
            route_id=route_id or None,
        )
        if early_route:
            request.route_id = early_route.id
            yield _sse({
                'type': 'route_resolved',
                'route_id': early_route.id,
                'route_name': early_route.name,
            })
            yield _sse({'type': 'status', 'status': 'generating', 'message': '正在生成报告，请稍候…'})

        def on_token(content: str):
            """LLM token 回调（同步，跨线程安全入队）"""
            if content:
                loop.call_soon_threadsafe(queue.put_nowait, ('token', content))

        def on_notice(content: str):
            """通知回调（模型降级等，同步入队）"""
            if content:
                loop.call_soon_threadsafe(queue.put_nowait, ('notice', content))

        def on_reasoning(content: str):
            """思考过程回调（仅开启思考的模型有内容）"""
            if content:
                loop.call_soon_threadsafe(queue.put_nowait, ('reasoning', content))

        task = asyncio.create_task(
            _service.generate_report(
                request=request,
                conversation_id=conversation_id,
                emit_event=None,
                on_token=on_token,
                on_notice=on_notice,
                on_reasoning=on_reasoning,
            )
        )

        try:
            # 先推流式 token，直到任务完成
            while not task.done():
                try:
                    kind, payload = await asyncio.wait_for(queue.get(), timeout=0.5)
                    sse = _make_stream_event(kind, payload)
                    if sse:
                        yield sse
                except asyncio.TimeoutError:
                    continue

            # 排空剩余 token / notice / reasoning
            while not queue.empty():
                kind, payload = queue.get_nowait()
                sse = _make_stream_event(kind, payload)
                if sse:
                    yield sse

            session, llm_output = task.result()
            _sessions[conversation_id] = session
            _persist_history(session)

            # 路由解析事件
            yield _sse({
                'type': 'route_resolved',
                'route_id': session.route_id,
                'route_name': session.route_name,
            })

            # 报告更新事件（含耗时）
            yield _sse({
                'type': 'report_updated',
                'session': session.dict(),
                'elapsed_seconds': session.elapsed_seconds,
            })

            # 追问
            if llm_output.questions:
                yield _sse({
                    'type': 'report_questions',
                    'questions': [q.dict() for q in llm_output.questions],
                })

            # 助手消息
            if llm_output.assistant_message:
                yield _sse({
                    'type': 'message',
                    'content': llm_output.assistant_message or '报告已生成。',
                })

            # 完成事件（含完整耗时）
            yield _sse({
                'type': 'done',
                'status': 'completed',
                'elapsed_seconds': session.elapsed_seconds,
            })

        except Exception as e:
            logger.error(f"生成报告失败: {e}", exc_info=True)
            if not task.done():
                task.cancel()
            yield _sse({'type': 'error', 'message': str(e)})

    return StreamingResponse(
        event_stream(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no",
        },
    )


@router.post("/sessions/{conversation_id}/iterate")
async def iterate_report(conversation_id: str, request: IterateRequest):
    """迭代报告（SSE流式）"""
    
    # 获取当前会话（内存缺失时从历史恢复，支持历史报告对话修改）
    session = _get_or_restore_session(conversation_id)
    if not session:
        raise HTTPException(status_code=404, detail="会话不存在")

    # 版本冲突保护：提交基于的版本落后于服务端当前版本时拒绝，防止并发/重复提交静默覆盖
    if request.base_version != session.current_version:
        raise HTTPException(
            status_code=409,
            detail=f"报告已被更新到 v{session.current_version}（本次提交基于 v{request.base_version}），请刷新报告后重试",
        )

    async def event_stream():
        """SSE事件流（真流式 token 推送 + 耗时记录）"""
        queue: asyncio.Queue = asyncio.Queue()
        loop = asyncio.get_event_loop()

        def on_token(content: str):
            if content:
                loop.call_soon_threadsafe(queue.put_nowait, ('token', content))

        def on_notice(content: str):
            if content:
                loop.call_soon_threadsafe(queue.put_nowait, ('notice', content))

        def on_reasoning(content: str):
            if content:
                loop.call_soon_threadsafe(queue.put_nowait, ('reasoning', content))

        task = asyncio.create_task(
            _service.iterate_report(
                request=request,
                current_session=session,
                emit_event=None,
                on_token=on_token,
                on_notice=on_notice,
                on_reasoning=on_reasoning,
            )
        )

        try:
            while not task.done():
                try:
                    kind, payload = await asyncio.wait_for(queue.get(), timeout=0.5)
                    sse = _make_stream_event(kind, payload)
                    if sse:
                        yield sse
                except asyncio.TimeoutError:
                    continue

            while not queue.empty():
                kind, payload = queue.get_nowait()
                sse = _make_stream_event(kind, payload)
                if sse:
                    yield sse

            updated_session, llm_output = task.result()
            _sessions[conversation_id] = updated_session
            _persist_history(updated_session)

            yield _sse({
                'type': 'report_updated',
                'session': updated_session.dict(),
                'elapsed_seconds': updated_session.elapsed_seconds,
            })

            if llm_output.questions:
                yield _sse({
                    'type': 'report_questions',
                    'questions': [q.dict() for q in llm_output.questions],
                })

            yield _sse({
                'type': 'message',
                'content': llm_output.assistant_message or '报告已更新。',
            })

            yield _sse({
                'type': 'done',
                'status': 'completed',
                'elapsed_seconds': updated_session.elapsed_seconds,
            })

        except Exception as e:
            logger.error(f"迭代报告失败: {e}", exc_info=True)
            if not task.done():
                task.cancel()
            yield _sse({'type': 'error', 'message': str(e)})

    return StreamingResponse(
        event_stream(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no",
        },
    )


@router.get("/sessions/{conversation_id}")
async def get_report_session(conversation_id: str):
    """获取当前报告"""
    session = _get_or_restore_session(conversation_id)
    if not session:
        raise HTTPException(status_code=404, detail="会话不存在")
    
    return ReportResponse(
        conversation_id=conversation_id,
        report=session,
    )


@router.post("/sessions/{conversation_id}/save")
async def save_report(conversation_id: str, request: SaveRequest):
    """手动保存报告"""
    session = _get_or_restore_session(conversation_id)
    if not session:
        raise HTTPException(status_code=404, detail="会话不存在")

    # 版本冲突保护：提交基于的版本落后于服务端当前版本时拒绝，防止并发/重复提交静默覆盖
    if request.base_version != session.current_version:
        raise HTTPException(
            status_code=409,
            detail=f"报告已被更新到 v{session.current_version}（本次提交基于 v{request.base_version}），请刷新报告后重试",
        )

    # 更新报告内容
    session.report.imaging_findings = request.imaging_findings
    session.report.imaging_diagnosis = request.imaging_diagnosis
    session.current_version += 1
    session.is_dirty = False
    
    from datetime import datetime
    session.updated_at = datetime.now().isoformat()
    
    _sessions[conversation_id] = session
    _persist_history(session)
    
    return {"success": True, "version": session.current_version}


@router.get("/routes")
async def list_routes():
    """列出所有可用的报告模板（含快捷输入标签）"""
    routes = _service.route_resolver.list_routes()
    return {
        "routes": [
            {
                "id": r.id,
                "name": r.name,
                "aliases": r.aliases,
                "quick_tags": r.quick_tags,
            }
            for r in routes
        ]
    }


@router.post("/sessions/{conversation_id}/export/word")
async def export_word(conversation_id: str):
    """导出报告为 Word 文档（.docx），返回文件流"""
    session = _get_or_restore_session(conversation_id)
    if not session:
        raise HTTPException(status_code=404, detail="会话不存在")

    try:
        from docx import Document
        from docx.shared import Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
    except ImportError:
        raise HTTPException(status_code=500, detail="服务端未安装 python-docx，无法导出 Word")

    from datetime import datetime

    doc = Document()

    # 全局中文字体
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    def set_cn_font(run, name="宋体"):
        run.font.name = name
        run._element.rPr.rFonts.set(qn("w:eastAsia"), name)

    # 报告头
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(f"{session.route_name} 影像报告")
    run.bold = True
    run.font.size = Pt(16)
    set_cn_font(run, "黑体")

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta.add_run(
        f"报告编号：{session.report_id}    生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}    版本：v{session.current_version}"
    )
    meta_run.font.size = Pt(9)
    set_cn_font(meta_run)

    def add_section(heading: str, text: str):
        h = doc.add_paragraph()
        h_run = h.add_run(heading)
        h_run.bold = True
        h_run.font.size = Pt(13)
        set_cn_font(h_run, "黑体")
        for line in (text or "").split("\n"):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(line)
            set_cn_font(r)

    add_section("影像所见", session.report.imaging_findings)
    add_section("影像诊断", session.report.imaging_diagnosis)

    # 免责声明
    note = doc.add_paragraph()
    note_run = note.add_run("本报告由AI辅助生成，仅供临床参考，最终诊断以医师审核签发为准。")
    note_run.font.size = Pt(8)
    set_cn_font(note_run)

    import io
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    filename = f"报告_{session.route_name}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
    from urllib.parse import quote
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )
